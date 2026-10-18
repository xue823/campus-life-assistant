# -*- coding: utf-8 -*-
"""为「黄淮校园生活助手」生成参赛答辩 PPT（16:9，嵌入真实界面截图）。"""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(BASE, "screenshots")
OUT = os.path.join(BASE, "黄淮校园生活助手-答辩PPT.pptx")
CROP_DIR = os.path.join(BASE, "_crop")

PRIMARY = RGBColor(0x18, 0x35, 0x62)
ACCENT = RGBColor(0x2A, 0x7F, 0xBC)
GOLD = RGBColor(0xB8, 0x7A, 0x12)
LIGHT = RGBColor(0xEE, 0xF4, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x30, 0x38, 0x41)
GREY = RGBColor(0x5A, 0x63, 0x6E)
CN, EN = "微软雅黑", "Segoe UI"

W, H = 13.333, 7.5
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


# --------------------------------------------------------------- 基础工具
def set_font(run, size, bold=False, color=INK, name=CN):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", name)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size, bold=False, color=INK, first=False,
         align=PP_ALIGN.LEFT, before=2, after=4, name=CN, spacing=1.15):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    p.line_spacing = spacing
    r = p.add_run()
    set_font(r, size, bold, color, name)
    r.text = text
    return p


def rect(slide, x, y, w, h, color, line=None, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def card(slide, x, y, w, h, fill=LIGHT, line=ACCENT, radius=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = rect(slide, x, y, w, h, fill, line, radius)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.14)
    tf.margin_bottom = Inches(0.1)
    return tf


def page_title(slide, title, sub=None):
    rect(slide, 0, 0, W, 0.14, PRIMARY)
    tf = textbox(slide, 0.6, 0.42, 12.1, 0.62)
    para(tf, title, 28, True, PRIMARY, first=True, before=0, after=0)
    if sub:
        tf2 = textbox(slide, 0.6, 1.06, 12.1, 0.36)
        para(tf2, sub, 14, False, GREY, first=True, before=0, after=0)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def phone(slide, img_path, x, y, h=4.55, crop_top=None, crop_ratio=None):
    """把截图裁成手机比例后贴入，带边框。"""
    if not os.path.isfile(img_path):
        return None
    im = Image.open(img_path).convert("RGB")
    w0, h0 = im.size
    if crop_top is not None:
        im = im.crop((0, 0, w0, min(h0, crop_top)))
    elif crop_ratio:
        im = im.crop((0, 0, w0, int(h0 * crop_ratio)))
    os.makedirs(CROP_DIR, exist_ok=True)
    out = os.path.join(CROP_DIR, "c_%s" % os.path.basename(img_path))
    im.save(out)
    ratio = im.width / im.height
    ph = Inches(h)
    pw = Inches(h * ratio)
    pic = slide.shapes.add_picture(out, Inches(x), Inches(y), width=pw, height=ph)
    pic.line.color.rgb = RGBColor(0xC9, 0xD5, 0xE4)
    pic.line.width = Pt(1)
    return pic


# =============================================================== 1 封面
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, WHITE)
rect(s, 0, 0, W, 0.3, PRIMARY)
rect(s, 0, H - 0.3, W, 0.3, PRIMARY)
rect(s, 0.75, 1.45, 0.14, 2.6, GOLD)

tf = textbox(s, 1.15, 1.35, 10.8, 0.4)
para(tf, "黄淮学院 · 校园生活助手", 15, True, ACCENT, first=True, before=0, after=0)
tf = textbox(s, 1.15, 1.85, 11.0, 1.0)
para(tf, "黄淮校园生活助手", 44, True, PRIMARY, first=True, before=0, after=2)
para(tf, "Campus Life Copilot · 一个把校园生活六件事装进手机的小工具", 17, False, GREY,
     before=0, after=0)

feats = ["考核安排 · 教学周自动过滤", "干饭推荐 · 理由可解释", "记账分析 · 每笔可改删",
         "同学名单 · 增删改查", "黄淮官方入口", "校园问答 · 离线可用"]
for i, f in enumerate(feats):
    col, row = i % 3, i // 3
    tf = textbox(s, 1.15 + col * 3.85, 3.55 + row * 0.62, 3.7, 0.5)
    para(tf, "● " + f, 13.5, False, INK, first=True, before=0, after=0)

tf = textbox(s, 1.15, 5.35, 11.0, 1.0)
para(tf, "Python 3 标准库 · 零第三方依赖 · 双击 start.bat 即运行 · 支持注册登录与数据增删改",
     13, True, GOLD, first=True, before=0, after=0)
para(tf, "GitHub：github.com/xue823/campus-life-assistant　　Gitee：gitee.com/x17993/campus-life-assistant",
     11.5, False, GREY, before=2, after=0)
notes(s, "开场：一句话说明这是给黄淮学院同学做的校园生活助手，六个模块、零依赖、双击就能跑。")

# =============================================================== 2 痛点
s = prs.slides.add_slide(BLANK)
page_title(s, "一、我们每天都在重复的六件小事", "问题不是“没有 App”，而是“信息散在六个地方”")
pains = [
    ("① 查考核安排", "教务系统入口难找，单双周、混合式教室记不清，跑错楼。"),
    ("② 到点吃什么", "食堂五六个、窗口几十个，排队十分钟，最后随便吃点。"),
    ("③ 钱花哪了", "月底发现钱没了，只记得“好像都吃了”，说不清构成。"),
    ("④ 找同学问事", "班委、舍友、同组同学的信息散在群聊里，问一次翻半天。"),
    ("⑤ 找不到系统入口", "教务、校历、图书馆、网络服务，每次都要搜或问学长。"),
    ("⑥ 校园事务问谁", "校园网认证、请假、补卡、快递……问一次要等半天回复。"),
]
for i, (t, d) in enumerate(pains):
    col, row = i % 3, i // 3
    tf = card(s, 0.6 + col * 4.12, 1.65 + row * 2.25, 3.85, 2.0,
              LIGHT if row == 0 else WHITE, ACCENT)
    para(tf, t, 15.5, True, PRIMARY, first=True, before=0, after=6)
    para(tf, d, 12.5, False, GREY, before=0, after=0, spacing=1.25)
notes(s, "六个痛点对应后面六个模块，一条一条对得上，别念完，挑三个讲。")

# =============================================================== 3 方案总览
s = prs.slides.add_slide(BLANK)
page_title(s, "二、方案：一个手机界面，六个模块", "打开即用、离线可跑、数据可改")
mods = [
    ("📅", "考核安排", "教学周引擎 · 倒计时 · 增删改"),
    ("🍚", "干饭推荐", "时段+口味+预算+排队 打分"),
    ("💰", "记账分析", "走势图 · 构成 · 省钱建议"),
    ("👥", "同学名单", "姓名/学号/联系方式 增删改查"),
    ("🏫", "黄淮信息", "10 个官方入口 · 联系电话"),
    ("🤖", "校园问答", "本地知识库 + 大模型降级"),
]
for i, (icon, name, desc) in enumerate(mods):
    col, row = i % 3, i // 3
    x, y = 0.6 + col * 4.12, 1.7 + row * 2.2
    tf = card(s, x, y, 3.85, 1.95, WHITE, ACCENT)
    para(tf, icon + "  " + name, 18, True, PRIMARY, first=True, before=0, after=5)
    para(tf, desc, 12.5, False, GREY, before=0, after=0, spacing=1.2)
    para(tf, "已实现 · 可演示", 11, True, GOLD, before=0, after=0)
tf = textbox(s, 0.6, 6.3, 12.1, 0.5)
para(tf, "共同底座：注册登录（PBKDF2 + Cookie 会话）· 数据按账号隔离 · 三类数据全部可修改可删除",
     13, False, ACCENT, first=True, before=0, after=0)
notes(s, "强调：六个模块不是六个页面堆砌，背后是同一套会话、数据隔离与增删改能力。")

# =============================================================== 4 考核安排
s = prs.slides.add_slide(BLANK)
page_title(s, "三、模块一：考核安排（教学周引擎）", "数据来自计科 2502B 考核安排表：24 条 / 12 门 / 27 学分")
phone(s, os.path.join(SHOTS, "02-schedule.png"), 0.7, 1.55, h=5.0, crop_top=1480)
points = [
    ("自动过滤单双周", "把「1-16 周」「单周」「双周」「3·7·11·14 周」解析成周次集合，按当前教学周过滤。"),
    ("下一场随时看得见", "首屏给出今天几场、下一场在哪、还有几天；正在考试会高亮。"),
    ("每条都能改能删", "修改 / 删除 / 一键清空，改完立刻重算学分与倒计时。"),
]
y = 1.6
for t, d in points:
    tf = card(s, 6.35, y, 6.35, 1.35, LIGHT, ACCENT)
    para(tf, t, 15, True, PRIMARY, first=True, before=0, after=4)
    para(tf, d, 12.5, False, GREY, before=0, after=0, spacing=1.2)
    y += 1.55
tf = textbox(s, 6.35, y + 0.05, 6.35, 0.8)
para(tf, "难点：单双周与指定周混排、同一门课不同周次不同考场 —— 解析器统一处理。",
     12.5, False, GOLD, first=True, before=0, after=0, spacing=1.2)
notes(s, "这是最有技术含量的一页：讲周次解析规则，再演示改一条、删一条，图表立即刷新。")

# =============================================================== 5 干饭推荐
s = prs.slides.add_slide(BLANK)
page_title(s, "四、模块二：干饭推荐（本地打分，可解释）", "食堂按黄淮实际：大榕树餐厅（南区）、北区第一/第二、西区小吃街")
phone(s, os.path.join(SHOTS, "03-meals.png"), 0.7, 1.55, h=5.0, crop_top=1450)
tf = card(s, 6.35, 1.6, 6.35, 1.75, LIGHT, ACCENT)
para(tf, "打分公式（0.4 秒内算完 34 道菜）", 15, True, PRIMARY, first=True, before=0, after=5)
para(tf, "时段匹配 ±22　口味命中 ×8　食堂偏好 ±30\n预算线性扣分　排队时间扣分　评分 ×12 加权",
     12.5, False, GREY, before=0, after=0, spacing=1.35)
for i, (t, d) in enumerate([
    ("每条都给理由", "「现在适合吃上午加餐；口味符合偏好；出餐快」——不是黑盒排序。"),
    ("三维筛选", "口味 / 食堂 / 单餐预算，偏好自动保存到账号。"),
    ("断网也能用", "纯本地计算，不依赖任何外卖平台接口。"),
]):
    tf = card(s, 6.35, 3.55 + i * 1.12, 6.35, 0.98, WHITE, ACCENT)
    para(tf, t, 13, True, PRIMARY, first=True, before=0, after=2)
    para(tf, d, 12, False, GREY, before=0, after=0, spacing=1.2)
notes(s, "重点讲“可解释”：每个推荐都带理由，评委能看懂算法在做什么。")

# =============================================================== 6 记账分析
s = prs.slides.add_slide(BLANK)
page_title(s, "五、模块三：记账分析", "两次点击记一笔，自动生成分析与建议")
phone(s, os.path.join(SHOTS, "04b-money-actions.png"), 0.7, 1.55, h=5.0, crop_top=1500)
for i, (t, d) in enumerate([
    ("自动分析", "今日 / 本月支出、14 天柱状走势（超支标橙）、消费构成、常去地点 TOP5。"),
    ("个性化建议", "按剩余预算与剩余天数算日均上限，指出花得最多的类别与地点。"),
    ("每笔可改可删", "点「改」表单预填原值，点「删」有确认；删除后图表与建议实时刷新。"),
    ("可导出", "一键导出当月 CSV，带 BOM 防止 Excel 中文乱码。"),
]):
    tf = card(s, 6.35, 1.6 + i * 1.24, 6.35, 1.08, LIGHT if i % 2 == 0 else WHITE, ACCENT)
    para(tf, t + "：", 13.5, True, PRIMARY, first=True, before=0, after=3)
    para(tf, d, 12, False, GREY, before=0, after=0, spacing=1.2)
notes(s, "现场加一笔 13.5 元，让评委看到数字和柱状图立刻变化。")

# =============================================================== 7 同学 + 黄淮
s = prs.slides.add_slide(BLANK)
page_title(s, "六、模块四、五：同学名单 与 黄淮信息", "一个管人，一个管入口")
phone(s, os.path.join(SHOTS, "07-students.png"), 0.75, 1.55, h=4.85, crop_top=1450)
phone(s, os.path.join(SHOTS, "08-campus.png"), 3.55, 1.55, h=4.85, crop_top=1450)
tf = card(s, 6.75, 1.6, 5.95, 2.2, LIGHT, ACCENT)
para(tf, "👥 同学名单", 15.5, True, PRIMARY, first=True, before=0, after=5)
para(tf, "姓名 / 学号 / 联系方式 / 身份 / 学院 / 备注，完整增删改查；\n"
         "数据按账号隔离，删除有确认弹窗；只有自己能看见。", 12.5, False, GREY,
     before=0, after=0, spacing=1.25)
tf = card(s, 6.75, 4.0, 5.95, 2.35, WHITE, GOLD)
para(tf, "🏫 黄淮信息（信息来自学校官网公开页面）", 15, True, GOLD, first=True, before=0, after=5)
para(tf, "学校概况：开源大道 76 号 · 18 个二级学院 · 61 个在招专业 · 校训「厚德、博学、笃行、自强」\n"
         "10 个官方入口：融合服务门户、教务管理系统、校历、网络服务、数字资源……\n"
         "联系电话：招生 0396-2853111 / 校园网服务中心 3728107、3728108",
     12, False, GREY, before=0, after=0, spacing=1.2)
notes(s, "这一页展示“落地”：不是泛泛演示数据，而是黄淮真实的入口、电话、食堂名称。")

# =============================================================== 8 问答
s = prs.slides.add_slide(BLANK)
page_title(s, "七、模块六：校园问答（双通道 + 降级）", "12 条本地知识库，覆盖校园网认证、注册登录、教务、校历、食堂、请假等")
phone(s, os.path.join(SHOTS, "05-chat.png"), 0.7, 1.55, h=5.0, crop_top=1450)
# 架构小图
tf = card(s, 6.35, 1.6, 6.35, 1.0, LIGHT, ACCENT)
para(tf, "① 关键词加权检索本地知识库", 13, True, PRIMARY, first=True, before=0, after=2)
para(tf, "毫秒级返回，零网络依赖，命中模糊时给“相近问题”。", 11.5, False, GREY,
     before=0, after=0)
tf = card(s, 6.35, 2.75, 6.35, 1.0, LIGHT, ACCENT)
para(tf, "② 配置 LLM_API_KEY 时调用大模型", 13, True, PRIMARY, first=True, before=0, after=2)
para(tf, "回答更自然、更长，适合解释型问题。", 11.5, False, GREY, before=0, after=0)
tf = card(s, 6.35, 3.9, 6.35, 1.0, WHITE, GOLD)
para(tf, "③ 大模型超时 / 失败 → 自动降级回本地库", 13, True, GOLD, first=True, before=0, after=2)
para(tf, "现场断网也不影响演示，这是刻意的工程取舍。", 11.5, False, GREY, before=0, after=0)
tf = textbox(s, 6.35, 5.15, 6.35, 1.3)
para(tf, "为什么这样做：比赛现场网络不可控。有网用大模型，没网照样答得出来；"
         "命中模糊时给“相近问题”，答不出时明确列出已知领域，不编造。",
     12.5, False, GREY, first=True, before=0, after=0, spacing=1.25)
notes(s, "演示时可以先问“校园网连不上怎么办”，再演示清空历史。强调降级设计。")

# =============================================================== 9 技术架构
s = prs.slides.add_slide(BLANK)
page_title(s, "八、技术架构：单文件后端 + 原生前端", "零第三方依赖，离线可运行，任何电脑拷过去双击即可")
layers = [
    ("前端（无框架、无 CDN）", "原生 JavaScript 单页应用 + 原生 CSS；桌面端手机外壳、移动端全屏；6 个标签页切换"),
    ("后端（Python 3 标准库）", "http.server 路由 + sqlite3；Cookie 会话、PBKDF2 加盐哈希、参数化 SQL、老库自动列迁移"),
    ("数据（本地 SQLite）", "users / courses / expenses / students / chat / prefs / sessions；首次启动自动建表并填充演示数据"),
    ("一键启动", "start.bat 自动识别随包 Python 运行时（36 MB），双击即跑并打开浏览器"),
]
y = 1.6
for i, (t, d) in enumerate(layers):
    tf = card(s, 0.6, y, 12.1, 1.12, WHITE if i % 2 else LIGHT, ACCENT)
    para(tf, t, 14.5, True, PRIMARY, first=True, before=0, after=3)
    para(tf, d, 12, False, GREY, before=0, after=0, spacing=1.2)
    y += 1.25
tf = textbox(s, 0.6, y + 0.05, 12.1, 0.9)
para(tf, "为什么不用框架：比赛现场不能保证有网、有环境；标准库 + 原生前端 = 零安装、零风险，"
         "而且每个接口、每行算法都能讲清楚。", 12.5, False, GOLD, first=True,
     before=0, after=0, spacing=1.2)
notes(s, "这页回答“为什么这么土的技术栈”：因为要现场跑、要能讲清楚、要零依赖。")

# =============================================================== 10 数据与亮点
s = prs.slides.add_slide(BLANK)
page_title(s, "九、工程数据与亮点", "可量化、可复现")
nums = [("18", "个接口", "注册登录 6 + 三类数据增删改查 9 + 问答/导出/重置"),
        ("1,900+", "行代码", "后端 1,000+ / 前端 600+ / 样式 350+"),
        ("27", "个上传文件", "代码 + 文档 + 15 张截图 + 封面图"),
        ("36 MB", "随包运行时", "免安装 Python，双击 start.bat 即跑")]
for i, (n, u, d) in enumerate(nums):
    x = 0.6 + i * 3.08
    tf = card(s, x, 1.6, 2.85, 1.6, LIGHT, ACCENT)
    para(tf, n, 26, True, PRIMARY, first=True, before=0, after=0)
    para(tf, u, 12.5, True, ACCENT, before=0, after=3)
    para(tf, d, 10.5, False, GREY, before=0, after=0, spacing=1.15)
high = [
    ("教学周引擎", "单双周 / 指定周自动过滤，把“这周到底上不上”彻底自动化"),
    ("推荐可解释", "每条推荐都写清理由，不是黑盒排序"),
    ("降级设计", "大模型不可用时自动回本地知识库，现场演示不会翻车"),
    ("双仓库 + 一键启动", "GitHub 与 Gitee 双份公开代码，随包运行时不装环境"),
]
y = 3.5
for i, (t, d) in enumerate(high):
    tf = card(s, 0.6, y, 12.1, 0.78, WHITE, GOLD)
    para(tf, "★ " + t + "　" + d, 12.5, False, GREY, first=True, before=0, after=0)
    y += 0.86
notes(s, "数字要背下来，问题“多少接口/多少行代码”能立刻答出。")

# =============================================================== 11 演示与链接
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, LIGHT)
rect(s, 0, 0, W, 0.16, PRIMARY)
rect(s, 0, H - 0.16, W, 0.16, PRIMARY)
tf = textbox(s, 0.6, 0.55, 12.1, 0.7)
para(tf, "十、现场演示与获取方式", 26, True, PRIMARY, first=True, before=0, after=0)

box = [
    ("演示（本机）", "http://127.0.0.1:8866/", "双击 start.bat → 自动打开浏览器"),
    ("演示（同一 Wi-Fi）", "http://10.121.96.97:8866/", "启动加 --host 0.0.0.0"),
    ("演示账号", "20230101 / 123456", "也可现场注册新账号体验"),
    ("代码仓库（GitHub）", "github.com/xue823/campus-life-assistant", "已公开，匿名可访问"),
    ("代码仓库（Gitee）", "gitee.com/x17993/campus-life-assistant", "国内访问更快"),
    ("封面图直链", "gitee.com/x17993/campus-life-assistant/raw/master/cover.png", "实测可显示"),
]
y = 1.45
for t, v, d in box:
    tf = card(s, 0.6, y, 12.1, 0.72, WHITE, ACCENT)
    para(tf, t + "：" + v + "　　" + d, 12, False, INK, first=True, before=0, after=0)
    y += 0.8

tf = textbox(s, 0.6, 6.55, 12.1, 0.7)
para(tf, "谢谢！欢迎提问，也欢迎反对我的设计取舍。", 17, True, GOLD, first=True,
     align=PP_ALIGN.CENTER, before=0, after=0)
notes(s, "结束语：留一个问题给自己——如果继续做，我会先做课表导入和考试提醒。")

prs.save(OUT)
print("saved:", OUT)
print("slides:", len(prs.slides._sldIdLst))
